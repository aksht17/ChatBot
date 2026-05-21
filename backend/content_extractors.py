import io
import os
import re
import tempfile
import zipfile
from urllib.parse import unquote, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup

from scraper import scrape_url

URL_RE = re.compile(r"https?://[^\s)\]]+")
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff")
AUDIO_VIDEO_EXTENSIONS = (
    ".mp3",
    ".wav",
    ".m4a",
    ".flac",
    ".ogg",
    ".opus",
    ".mp4",
    ".mov",
    ".avi",
    ".mkv",
    ".webm",
)
ARCHIVE_EXTENSIONS = (".zip",)
DOWNLOAD_EXTENSIONS = (
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".xls",
    ".csv",
    ".html",
    ".htm",
    ".txt",
    ".md",
) + IMAGE_EXTENSIONS + AUDIO_VIDEO_EXTENSIONS + ARCHIVE_EXTENSIONS

_WHISPER_CACHE = {}


def _normalize_host(url):
    host = urlparse(url).hostname or ""
    return host.lower().lstrip(".")


def _domain_matches(host, domain):
    host = (host or "").lower().lstrip(".")
    domain = (domain or "").lower().lstrip(".")
    if not host or not domain:
        return False
    return host == domain or host.endswith(f".{domain}")


def _is_url_allowed(url, allowed_domains=None, blocked_domains=None):
    host = _normalize_host(url)
    if not host:
        return False

    blocked = blocked_domains or []
    if any(_domain_matches(host, domain) for domain in blocked):
        return False

    allowed = allowed_domains or []
    if allowed:
        return any(_domain_matches(host, domain) for domain in allowed)
    return True


def _extract_urls(text, max_urls):
    urls = URL_RE.findall(text)
    if max_urls is None:
        return urls
    if isinstance(max_urls, int) and max_urls > 0:
        return urls[:max_urls]
    return []


def _strip_blocked_urls_from_text(text, blocked_domains=None):
    blocked = blocked_domains or []
    if not blocked:
        return text

    cleaned = text
    for url in URL_RE.findall(text):
        if not _is_url_allowed(url, allowed_domains=None, blocked_domains=blocked):
            cleaned = cleaned.replace(url, " ")
    return cleaned


def _is_downloadable_url(url):
    path = urlparse(url).path.lower()
    return path.endswith(DOWNLOAD_EXTENSIONS) or "/cdn/view/" in path


def _filename_from_url(url, content_type):
    path_name = os.path.basename(unquote(urlparse(url).path))
    if path_name and "." in path_name:
        return path_name

    if "pdf" in content_type:
        return "downloaded.pdf"
    if "csv" in content_type:
        return "downloaded.csv"
    if "spreadsheet" in content_type or "excel" in content_type:
        return "downloaded.xlsx"
    if "presentation" in content_type:
        return "downloaded.pptx"
    if "wordprocessing" in content_type or "msword" in content_type:
        return "downloaded.docx"
    if "html" in content_type:
        return "downloaded.html"
    if "json" in content_type:
        return "downloaded.json"
    return path_name or "downloaded.txt"


def _looks_like_download_url(url):
    path = urlparse(url).path.lower()
    return path.endswith(DOWNLOAD_EXTENSIONS) or "/cdn/view/" in path


def _documents_from_url(
    url,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl,
    allowed_domains=None,
    blocked_domains=None,
):
    if not _is_url_allowed(url, allowed_domains, blocked_domains):
        print(f"Skipped URL outside crawl policy: {url}")
        return []

    if not _looks_like_download_url(url):
        if not allow_external_web_crawl:
            return []
        return scrape_url(
            url,
            depth=follow_links_depth,
            allowed_domains=allowed_domains,
            blocked_domains=blocked_domains,
        )

    try:
        response = requests.get(
            url,
            timeout=30,
            headers={"User-Agent": "Mozilla/5.0"},
        )
        response.raise_for_status()
    except Exception as exc:
        print(f"Failed download {url}: {exc}")
        return []

    content_type = response.headers.get("content-type", "").lower()
    name = _filename_from_url(url, content_type)
    documents = extract_from_bytes(
        name,
        response.content,
        follow_links_depth=follow_links_depth,
        follow_urls_in_text=False,
        max_urls=max_urls,
        allow_external_web_crawl=allow_external_web_crawl,
        allowed_domains=allowed_domains,
        blocked_domains=blocked_domains,
    )
    for document in documents:
        document["source"] = url
    print(f"Extracted {len(documents)} documents from linked file: {url}")
    return documents


def documents_from_text(
    text,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    documents = []
    cleaned = _strip_blocked_urls_from_text(text, blocked_domains).strip()
    if cleaned:
        documents.append({"source": source, "content": cleaned})

    if follow_urls_in_text and cleaned:
        for url in _extract_urls(cleaned, max_urls):
            documents += _documents_from_url(
                url,
                follow_links_depth,
                follow_urls_in_text,
                max_urls,
                allow_external_web_crawl,
                allowed_domains,
                blocked_domains,
            )

    return documents


def documents_from_dataframe(
    df,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    documents = []
    for _, row in df.iterrows():
        parts = []
        for col in df.columns:
            parts.append(f"{col}: {row[col]}")
        row_text = " | ".join(parts)
        documents += documents_from_text(
            row_text,
            source,
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )
    return documents


def _documents_from_html(
    content,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    soup = BeautifulSoup(content, "html.parser")
    text = soup.get_text(separator=" ", strip=True)
    documents = documents_from_text(
        text,
        source,
        follow_links_depth,
        follow_urls_in_text,
        max_urls,
        allow_external_web_crawl,
        allowed_domains,
        blocked_domains,
    )
    if follow_links_depth > 0:
        for a in soup.find_all("a", href=True):
            link = a["href"]
            if (
                link.startswith("http")
                and allow_external_web_crawl
                and _is_url_allowed(link, allowed_domains, blocked_domains)
            ):
                documents += scrape_url(
                    link,
                    depth=follow_links_depth - 1,
                    allowed_domains=allowed_domains,
                    blocked_domains=blocked_domains,
                )
    return documents


def _documents_from_pdf(
    content,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    try:
        from pypdf import PdfReader
    except Exception:
        print("Skipping PDF: missing pypdf")
        return []

    reader = PdfReader(io.BytesIO(content))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    return documents_from_text(
        "\n".join(pages),
        source,
        follow_links_depth,
        follow_urls_in_text,
        max_urls,
        allow_external_web_crawl,
        allowed_domains,
        blocked_domains,
    )


def _documents_from_docx(
    content,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    try:
        import docx
    except Exception:
        print("Skipping DOCX: missing python-docx")
        return []

    doc = docx.Document(io.BytesIO(content))
    text = "\n".join([p.text for p in doc.paragraphs])
    return documents_from_text(
        text,
        source,
        follow_links_depth,
        follow_urls_in_text,
        max_urls,
        allow_external_web_crawl,
        allowed_domains,
        blocked_domains,
    )


def _documents_from_pptx(
    content,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    try:
        from pptx import Presentation
    except Exception:
        print("Skipping PPTX: missing python-pptx")
        return []

    pres = Presentation(io.BytesIO(content))
    text_chunks = []
    for slide in pres.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                text_chunks.append(text)
    return documents_from_text(
        "\n".join(text_chunks),
        source,
        follow_links_depth,
        follow_urls_in_text,
        max_urls,
        allow_external_web_crawl,
        allowed_domains,
        blocked_domains,
    )


def _documents_from_image(
    content,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    try:
        from PIL import Image
        import pytesseract
    except Exception:
        print("Skipping image OCR: missing pillow/pytesseract")
        return []

    img = Image.open(io.BytesIO(content))
    text = pytesseract.image_to_string(img)
    return documents_from_text(
        text,
        source,
        follow_links_depth,
        follow_urls_in_text,
        max_urls,
        allow_external_web_crawl,
        allowed_domains,
        blocked_domains,
    )


def _get_whisper_model(model_name, device, compute_type):
    key = (model_name, device, compute_type)
    if key in _WHISPER_CACHE:
        return _WHISPER_CACHE[key]
    from faster_whisper import WhisperModel

    model = WhisperModel(model_name, device=device, compute_type=compute_type)
    _WHISPER_CACHE[key] = model
    return model


def _to_wav(input_path):
    try:
        import ffmpeg
    except Exception:
        print("Skipping audio/video: missing ffmpeg-python")
        return None

    fd, out_path = tempfile.mkstemp(suffix=".wav")
    os.close(fd)
    try:
        (
            ffmpeg.input(input_path)
            .output(out_path, ac=1, ar=16000)
            .overwrite_output()
            .run(quiet=True)
        )
    except Exception:
        try:
            os.remove(out_path)
        except OSError:
            pass
        return None
    return out_path


def _documents_from_media(
    content,
    source,
    follow_links_depth,
    follow_urls_in_text,
    max_urls,
    model_name,
    device,
    compute_type,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    try:
        model = _get_whisper_model(model_name, device, compute_type)
    except Exception:
        print("Skipping audio/video: missing faster-whisper")
        return []

    fd, input_path = tempfile.mkstemp(suffix=".media")
    os.close(fd)
    with open(input_path, "wb") as handle:
        handle.write(content)

    wav_path = _to_wav(input_path)
    try:
        if not wav_path:
            return []
        segments, _ = model.transcribe(wav_path)
        text = " ".join([seg.text.strip() for seg in segments if seg.text])
        return documents_from_text(
            text,
            source,
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )
    finally:
        try:
            os.remove(input_path)
        except OSError:
            pass
        if wav_path:
            try:
                os.remove(wav_path)
            except OSError:
                pass


def extract_from_bytes(
    name,
    content,
    follow_links_depth=1,
    follow_urls_in_text=True,
    max_urls=20,
    enable_ocr=True,
    enable_audio_video=True,
    whisper_model="base",
    whisper_device="cpu",
    whisper_compute_type="int8",
    archive_depth=1,
    allow_external_web_crawl=True,
    allowed_domains=None,
    blocked_domains=None,
):
    lower = name.lower()

    if lower.endswith(ARCHIVE_EXTENSIONS) and archive_depth > 0:
        documents = []
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            for inner in zf.namelist():
                if inner.endswith("/"):
                    continue
                data = zf.read(inner)
                documents += extract_from_bytes(
                    inner,
                    data,
                    follow_links_depth,
                    follow_urls_in_text,
                    max_urls,
                    enable_ocr,
                    enable_audio_video,
                    whisper_model,
                    whisper_device,
                    whisper_compute_type,
                    archive_depth=archive_depth - 1,
                    allow_external_web_crawl=allow_external_web_crawl,
                    allowed_domains=allowed_domains,
                    blocked_domains=blocked_domains,
                )
        return documents

    if lower.endswith((".xlsx", ".xls")):
        df = pd.read_excel(io.BytesIO(content))
        return documents_from_dataframe(
            df,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    if lower.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(content))
        return documents_from_dataframe(
            df,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    if lower.endswith((".html", ".htm")):
        return _documents_from_html(
            content,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    if lower.endswith(".pdf"):
        return _documents_from_pdf(
            content,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    if lower.endswith(".docx"):
        return _documents_from_docx(
            content,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    if lower.endswith(".pptx"):
        return _documents_from_pptx(
            content,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    if lower.endswith(IMAGE_EXTENSIONS):
        if not enable_ocr:
            return []
        return _documents_from_image(
            content,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    if lower.endswith(AUDIO_VIDEO_EXTENSIONS):
        if not enable_audio_video:
            return []
        return _documents_from_media(
            content,
            f"file:{name}",
            follow_links_depth,
            follow_urls_in_text,
            max_urls,
            whisper_model,
            whisper_device,
            whisper_compute_type,
            allow_external_web_crawl,
            allowed_domains,
            blocked_domains,
        )

    try:
        text = content.decode(errors="ignore")
    except Exception:
        text = ""
    return documents_from_text(
        text,
        f"file:{name}",
        follow_links_depth,
        follow_urls_in_text,
        max_urls,
        allow_external_web_crawl,
        allowed_domains,
        blocked_domains,
    )
